import openpyxl
import re
from openpyxl.styles import PatternFill
from datetime import datetime
from rapidfuzz import process, fuzz
from pathlib import Path

# Green/red fill styles
green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
grey_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

def normalize_indicator_name(name):
    if not name:
        return ""

    # 1. Remove trailing parentheticals like "(Apr)", "(YoY)", etc.
    name = re.sub(r"\s*(\([^)]*\))*\s*$", "", name)

    # 2. Remove punctuation but keep alphanumeric + spaces
    name = re.sub(r"[^a-zA-Z0-9\s]", "", name)

    # 3. Remove low-signal stop words (DO NOT include 'income', 'avg', 'wage', etc.)
    stop_words = r"\b(yoy|mom|index|total|report|data|figures)\b"
    name = re.sub(stop_words, "", name, flags=re.IGNORECASE)

    # 4. Normalize spacing and case
    return name.strip().lower()
    
    # Get best fuzzy match using RapidFuzz
def find_best_match(target, options):
    result = process.extractOne(
        target,
        options,
        scorer=fuzz.token_sort_ratio,
        score_cutoff=80  # Adjust cutoff as needed
    )
    if result:
        print(f"🔍 Matched '{target}' → '{result[0]}' with score {result[1]}")
        return result[0]
    return None

# Convert date to Excel column (Jan–Dec → D–O)
def get_month_column(date_str):
    dt = datetime.strptime(date_str, "%Y/%m/%d %H:%M:%S")
    return chr(ord("D") + dt.month - 1)

def write_to_excel(data, country_code):
    path = Path(__file__).parent / "Economic_Analysis_Draft.xlsx"
    wb = openpyxl.load_workbook(path)

    try:
        staging_ws = wb["Forecast Data"]
    except KeyError:
        print("❌ Worksheet 'Forecast Data' not found.")
        return

    target_ws = wb[country_code]

    # Build map of normalized names → row numbers
    sheet_name_map = {
        normalize_indicator_name(target_ws.cell(row=row, column=2).value): row
        for row in range(2, target_ws.max_row + 1)
        if target_ws.cell(row=row, column=2).value
    }

    for row_data in data:
        if not row_data["actual"]:
            print(f"⏩ Skipping unreleased indicator: {row_data['name']} at {row_data['date']}")
            continue

        normalized_name = normalize_indicator_name(row_data["name"])
        best_match = find_best_match(normalized_name, list(sheet_name_map.keys()))

        if best_match:
            matched_row = sheet_name_map[best_match]
            col_letter = get_month_column(row_data["date"])
            cell = target_ws[f"{col_letter}{matched_row}"]

            # Get style info
            actual_class = row_data.get("actual_class", "")

            # Clean actual value
            clean_actual = row_data["actual"].strip()
            clean_actual = re.sub(r'[\s]*[BMKbmk]$', '', clean_actual)

            try:
                if '%' in clean_actual:
                    cell.value = float(clean_actual.replace('%', '').replace(',', '')) / 100
                else:
                    cell.value = float(clean_actual.replace(',', ''))
            except ValueError:
                cell.value = clean_actual  # fallback to string

           # Apply conditional formatting
            forecast_value = row_data.get("forecast", "").strip()

            # Try to compare actual and forecast numerically
            try:
                numeric_actual = float(clean_actual.replace('%', '').replace(',', '')) / 100 if '%' in clean_actual else float(clean_actual.replace(',', ''))
                numeric_forecast = float(forecast_value.replace('%', '').replace(',', '')) / 100 if '%' in forecast_value else float(forecast_value.replace(',', ''))

                if abs(numeric_actual - numeric_forecast) < 1e-6:  # allow small float difference
                    cell.fill = grey_fill
                elif "greenFont" in actual_class:
                    cell.fill = green_fill
                elif "redFont" in actual_class:
                    cell.fill = red_fill
            except ValueError:
                # Fallback to class-based color if conversion fails
                if "greenFont" in actual_class:
                    cell.fill = green_fill
                elif "redFont" in actual_class:
                    cell.fill = red_fill

        else:
            print(f"⚠️ Could not find row for: {row_data['name']}")

    wb.save(path)
    print(f"✅ Data written to '{path}'")
